#!/usr/bin/env python3
"""
clockify_fill.py — Clockify Time Entry Automation Tool

Reads tickets and a Clockify CSV export, generates a per-day plan
filling remaining time (up to 8h/day, window 12:00–20:00), shows
the plan for approval, then creates entries via the Clockify API.
"""

import argparse
import csv
import json
import sys
import time
from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo

try:
    import requests
except ImportError:
    print("Error: 'requests' library not found. Install it with: pip install requests")
    sys.exit(1)

# Ensure Unicode output works on Windows (UTF-8 console or file redirect)
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

# ─── ANSI colours ─────────────────────────────────────────────────────────────
GREY  = "\033[90m"
GREEN = "\033[32m"
RED   = "\033[31m"
BOLD  = "\033[1m"
RESET = "\033[0m"

CLOCKIFY_API_BASE  = "https://api.clockify.me/api/v1"
WORK_DAY_MINUTES   = 8 * 60   # 480
MIN_ENTRY_MINUTES  = 30
RATE_LIMIT_DELAY_S = 0.2


# ─── Data structures ──────────────────────────────────────────────────────────

class Ticket:
    def __init__(self, key: str, title: str, project_id: str,
                 date_from: date | None = None, date_to: date | None = None):
        self.key        = key.strip()
        self.title      = title.strip()
        self.project_id = project_id.strip()
        self.date_from  = date_from   # None means no lower bound
        self.date_to    = date_to     # None means no upper bound

    @property
    def description(self) -> str:
        return f"{self.key} {self.title}"


class TimeSlot:
    """A time interval on a specific date (naive local time)."""

    def __init__(self, day: date, start: datetime, end: datetime, *,
                 description: str = "", project_id: str = "", existing: bool = False):
        self.day         = day
        self.start       = start
        self.end         = end
        self.description = description
        self.project_id  = project_id
        self.existing    = existing

    @property
    def duration_minutes(self) -> int:
        return max(0, int((self.end - self.start).total_seconds() / 60))

    @property
    def duration_hours(self) -> float:
        return self.duration_minutes / 60


# ─── Config / file parsing ────────────────────────────────────────────────────

def _strip_json_comments(text: str) -> str:
    """Remove // line-comments that are not inside JSON string literals."""
    result    = []
    in_string = False
    i         = 0
    while i < len(text):
        ch = text[i]
        if ch == '\\' and in_string:
            # Escaped character — keep both chars as-is
            result.append(ch)
            i += 1
            if i < len(text):
                result.append(text[i])
                i += 1
            continue
        if ch == '"':
            in_string = not in_string
            result.append(ch)
            i += 1
            continue
        if not in_string and text[i:i+2] == '//':
            # Skip until end of line
            while i < len(text) and text[i] != '\n':
                i += 1
            continue
        result.append(ch)
        i += 1
    return ''.join(result)


def load_config(path: str) -> dict:
    with open(path, encoding="utf-8") as fh:
        raw = fh.read()
    return json.loads(_strip_json_comments(raw))


def _parse_ticket_date(s: str, lineno: int, field: str) -> date | None:
    s = s.strip()
    if not s:
        return None
    try:
        return datetime.strptime(s, "%Y-%m-%d").date()
    except ValueError:
        print(f"Warning: line {lineno} — invalid {field} date '{s}' (expected YYYY-MM-DD), ignoring.")
        return None


def load_tickets(path: str) -> list[Ticket]:
    tickets = []
    with open(path, encoding="utf-8") as fh:
        for lineno, line in enumerate(fh, 1):
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            parts = [p.strip() for p in line.split("|")]
            if len(parts) < 3:
                print(f"Warning: line {lineno} in tickets file has "
                      f"{len(parts)} field(s) (expected 3–5), skipping.")
                continue
            date_from = _parse_ticket_date(parts[3], lineno, "date_from") if len(parts) > 3 else None
            date_to   = _parse_ticket_date(parts[4], lineno, "date_to")   if len(parts) > 4 else None
            tickets.append(Ticket(parts[0], parts[1], parts[2], date_from, date_to))
    return tickets


# ─── CSV parsing ──────────────────────────────────────────────────────────────

_DATE_FORMATS = ("%m/%d/%Y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y")
_TIME_FORMATS = ("%H:%M:%S", "%H:%M", "%I:%M:%S %p", "%I:%M %p")


def _parse_date(s: str) -> date | None:
    s = s.strip()
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def _parse_time(s: str, ref_date: date) -> datetime | None:
    s = s.strip()
    for fmt in _TIME_FORMATS:
        try:
            return datetime.combine(ref_date, datetime.strptime(s, fmt).time())
        except ValueError:
            pass
    return None


def _window_bounds(day: date, work_start: str, work_end: str) -> tuple[datetime, datetime]:
    ws_h, ws_m = map(int, work_start.split(":"))
    we_h, we_m = map(int, work_end.split(":"))
    return (
        datetime(day.year, day.month, day.day, ws_h, ws_m),
        datetime(day.year, day.month, day.day, we_h, we_m),
    )


def parse_clockify_csv(path: str, target_month: str,
                       work_start: str, work_end: str) -> dict[date, list[TimeSlot]]:
    """Return {date: [TimeSlot, ...]} for existing entries clipped to the work window."""
    year, month = map(int, target_month.split("-"))
    entries: dict[date, list[TimeSlot]] = {}

    try:
        with open(path, newline="", encoding="utf-8-sig") as fh:
            reader = csv.DictReader(fh)
            for row in reader:
                start_day = _parse_date(row.get("Start Date", ""))
                if start_day is None:
                    continue
                if start_day.year != year or start_day.month != month:
                    continue

                start_dt = _parse_time(row.get("Start Time", ""), start_day)
                if start_dt is None:
                    continue

                end_day_str = (row.get("End Date") or "").strip()
                end_day     = _parse_date(end_day_str) if end_day_str else start_day
                if end_day is None:
                    end_day = start_day
                end_dt = _parse_time(row.get("End Time", ""), end_day)
                if end_dt is None:
                    continue

                win_start, win_end = _window_bounds(start_day, work_start, work_end)
                clipped_start = max(start_dt, win_start)
                clipped_end   = min(end_dt,   win_end)
                if clipped_end <= clipped_start:
                    continue

                desc    = (row.get("Description") or row.get("Task") or "").strip()
                project = (row.get("Project") or "").strip()
                label   = f"{desc} ({project})" if project and desc else desc or project or "entry"

                slot = TimeSlot(start_day, clipped_start, clipped_end,
                                description=label, existing=True)
                entries.setdefault(start_day, []).append(slot)

    except FileNotFoundError:
        print(f"Note: CSV report not found at '{path}'. Assuming no existing entries.")

    for d in entries:
        entries[d].sort(key=lambda s: s.start)
    return entries


# ─── Working-day helpers ──────────────────────────────────────────────────────

def get_working_days(target_month: str, skip_weekends: bool) -> list[date]:
    year, month = map(int, target_month.split("-"))
    days = []
    d = date(year, month, 1)
    while d.month == month:
        if not skip_weekends or d.weekday() < 5:
            days.append(d)
        d += timedelta(days=1)
    return days


def day_logged_minutes(existing: list[TimeSlot]) -> int:
    return sum(s.duration_minutes for s in existing)


def compute_free_slots(day: date, existing: list[TimeSlot],
                       work_start: str, work_end: str) -> list[TimeSlot]:
    """Return gaps within the work window not covered by existing entries."""
    win_start, win_end = _window_bounds(day, work_start, work_end)
    occupied = sorted(existing, key=lambda s: s.start)
    free   = []
    cursor = win_start
    for slot in occupied:
        if slot.start > cursor:
            free.append(TimeSlot(day, cursor, slot.start))
        cursor = max(cursor, slot.end)
    if cursor < win_end:
        free.append(TimeSlot(day, cursor, win_end))
    return free


# ─── Distribution algorithm ───────────────────────────────────────────────────

def _ticket_distance(ticket: Ticket, day: date) -> int:
    """Days between `day` and the ticket's date range (0 if covered)."""
    if ticket.date_from and day < ticket.date_from:
        return (ticket.date_from - day).days
    if ticket.date_to and day > ticket.date_to:
        return (day - ticket.date_to).days
    return 0


def get_active_tickets_for_day(day: date, tickets: list[Ticket]) -> list[Ticket]:
    """
    Returns tickets whose date range covers `day`.
    Tickets with no date range always cover every day.
    If no ticket covers the day, returns the ticket(s) with the closest range.
    """
    if not tickets:
        return []
    covering = [t for t in tickets if _ticket_distance(t, day) == 0]
    if covering:
        return covering
    # Fall back to closest ticket(s) by date proximity
    min_dist = min(_ticket_distance(t, day) for t in tickets)
    return [t for t in tickets if _ticket_distance(t, day) == min_dist]


def distribute_tickets(working_days: list[date],
                       existing_by_day: dict[date, list[TimeSlot]],
                       tickets: list[Ticket],
                       work_start: str, work_end: str) -> list[TimeSlot]:
    """
    For each working day, find the active tickets (those whose date range covers
    the day, or the closest ticket(s) if none cover it), then split the day's
    free time evenly among them.
    """
    if not tickets:
        return []

    new_entries: list[TimeSlot] = []

    for day in working_days:
        existing     = existing_by_day.get(day, [])
        logged       = day_logged_minutes(existing)
        still_needed = WORK_DAY_MINUTES - logged
        if still_needed <= 0:
            continue

        free_slots  = compute_free_slots(day, existing, work_start, work_end)
        free_in_day = sum(s.duration_minutes for s in free_slots)
        usable      = min(still_needed, free_in_day)

        if usable < MIN_ENTRY_MINUTES:
            continue

        active = get_active_tickets_for_day(day, tickets)
        # Budget per ticket for this day (even split)
        budget_per_ticket = usable / len(active)
        ticket_budget  = {t.key: budget_per_ticket for t in active}
        ticket_by_key  = {t.key: t for t in active}
        ticket_keys    = [t.key for t in active]
        ticket_idx     = 0
        day_remaining  = still_needed

        for free_slot in free_slots:
            slot_remaining = free_slot.duration_minutes
            cursor         = free_slot.start

            while slot_remaining >= MIN_ENTRY_MINUTES and ticket_idx < len(ticket_keys):
                tkey   = ticket_keys[ticket_idx]
                ticket = ticket_by_key[tkey]
                t_rem  = ticket_budget[tkey]

                if t_rem < MIN_ENTRY_MINUTES:
                    ticket_idx += 1
                    continue

                assign = min(slot_remaining, t_rem, day_remaining)

                if assign < MIN_ENTRY_MINUTES:
                    break

                # Absorb tiny leftovers to avoid orphan sub-30-min gaps
                leftover = slot_remaining - assign
                if 0 < leftover < MIN_ENTRY_MINUTES:
                    assign = slot_remaining

                entry_end = cursor + timedelta(minutes=assign)
                new_entries.append(TimeSlot(
                    day, cursor, entry_end,
                    description=ticket.description,
                    project_id=ticket.project_id,
                    existing=False,
                ))

                ticket_budget[tkey] -= assign
                slot_remaining      -= assign
                day_remaining       -= assign
                cursor               = entry_end

                if ticket_budget[tkey] < MIN_ENTRY_MINUTES:
                    ticket_idx += 1

            if day_remaining < MIN_ENTRY_MINUTES:
                break

    return new_entries


# ─── Display ──────────────────────────────────────────────────────────────────

def _fmt_time(dt: datetime) -> str:
    return dt.strftime("%H:%M")


def _fmt_dur(minutes: int) -> str:
    h, m = divmod(minutes, 60)
    return f"{h}h{m:02d}m" if m else f"{h}h"


def print_plan(working_days: list[date],
               existing_by_day: dict[date, list[TimeSlot]],
               new_entries: list[TimeSlot]):
    new_by_day: dict[date, list[TimeSlot]] = {}
    for e in new_entries:
        new_by_day.setdefault(e.day, []).append(e)

    W = 80
    print(f"\n{'─' * W}")
    print(f"{'DATE':<13} {'START':<7} {'END':<7} {'DUR':<8}  DESCRIPTION")
    print(f"{'─' * W}")

    days_complete = 0

    for day in working_days:
        existing = existing_by_day.get(day, [])
        new      = new_by_day.get(day, [])
        if not existing and not new:
            continue

        all_entries  = sorted(existing + new, key=lambda s: s.start)
        total_min    = sum(s.duration_minutes for s in all_entries)
        is_complete  = total_min >= WORK_DAY_MINUTES - 1
        if is_complete:
            days_complete += 1

        tag = "  ✓" if is_complete else f"  ({_fmt_dur(total_min)}/8h)"
        print(f"\n{BOLD}{day.strftime('%Y-%m-%d (%a)')}{tag}{RESET}")

        for slot in all_entries:
            dur = _fmt_dur(slot.duration_minutes)
            if slot.existing:
                print(f"  {GREY}{_fmt_time(slot.start):<7} {_fmt_time(slot.end):<7}"
                      f" {dur:<8}  [existing] {slot.description}{RESET}")
            else:
                print(f"  {_fmt_time(slot.start):<7} {_fmt_time(slot.end):<7}"
                      f" {dur:<8}  {slot.description}")

    print(f"\n{'─' * W}")

    # Per-ticket summary
    ticket_minutes: dict[str, int] = {}
    for e in new_entries:
        ticket_minutes[e.description] = ticket_minutes.get(e.description, 0) + e.duration_minutes

    days_already_full = sum(
        1 for d in working_days
        if day_logged_minutes(existing_by_day.get(d, [])) >= WORK_DAY_MINUTES - 1
    )
    total_new_min = sum(e.duration_minutes for e in new_entries)

    print(f"\nSummary:")
    print(f"  Working days in month : {len(working_days)}")
    print(f"  Days already complete : {days_already_full}")
    print(f"  Days with new entries : {len(new_by_day)}")
    print(f"  New entries to create : {len(new_entries)}")
    print(f"  Total hours to add    : {_fmt_dur(total_new_min)}")

    if ticket_minutes:
        print(f"\n  Per-ticket allocation:")
        for desc, mins in sorted(ticket_minutes.items()):
            print(f"    {_fmt_dur(mins):>8}  {desc}")
    print()


# ─── Clockify API ─────────────────────────────────────────────────────────────

def _local_to_utc(dt: datetime, tz: ZoneInfo) -> datetime:
    return dt.replace(tzinfo=tz).astimezone(ZoneInfo("UTC"))


def _create_entry(session: requests.Session, workspace_id: str,
                  slot: TimeSlot, tz: ZoneInfo) -> tuple[bool, str]:
    url = f"{CLOCKIFY_API_BASE}/workspaces/{workspace_id}/time-entries"
    payload = {
        "description": slot.description,
        "projectId":   slot.project_id,
        "start":       _local_to_utc(slot.start, tz).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "end":         _local_to_utc(slot.end,   tz).strftime("%Y-%m-%dT%H:%M:%SZ"),
    }
    resp = session.post(url, json=payload, timeout=30)
    if resp.status_code in (200, 201):
        return True, ""
    try:
        msg = resp.json().get("message", resp.text)
    except Exception:
        msg = resp.text
    return False, f"HTTP {resp.status_code}: {msg}"


def push_entries(new_entries: list[TimeSlot],
                 api_key: str, workspace_id: str, tz: ZoneInfo):
    session = requests.Session()
    session.headers.update({
        "X-Api-Key":    api_key,
        "Content-Type": "application/json",
    })

    created = failed = 0
    for slot in new_entries:
        label = (f"{slot.description} | "
                 f"{slot.day} {_fmt_time(slot.start)}–{_fmt_time(slot.end)}")
        try:
            ok, err = _create_entry(session, workspace_id, slot, tz)
        except requests.RequestException as exc:
            ok, err = False, str(exc)

        if ok:
            print(f"{GREEN}✓{RESET} {label}")
            created += 1
        else:
            print(f"{RED}✗{RESET} {label} | Error: {err}")
            failed += 1

        time.sleep(RATE_LIMIT_DELAY_S)

    print(f"\nCreated: {created} entries | Failed: {failed}")


# ─── Plan serialisation ───────────────────────────────────────────────────────

def _build_plan_json(working_days: list[date],
                     existing_by_day: dict[date, list[TimeSlot]],
                     new_entries: list[TimeSlot],
                     tz: ZoneInfo) -> list[dict]:
    """Build the plan.json array ordered chronologically by day and start time."""
    new_by_day: dict[date, list[TimeSlot]] = {}
    for slot in new_entries:
        new_by_day.setdefault(slot.day, []).append(slot)

    plan = []
    for day in working_days:
        all_slots = sorted(
            existing_by_day.get(day, []) + new_by_day.get(day, []),
            key=lambda s: s.start,
        )
        for slot in all_slots:
            if slot.existing:
                plan.append({
                    "description": slot.description,
                    "projectId":   "",
                    "start":       _local_to_utc(slot.start, tz).strftime("%Y-%m-%dT%H:%M:%SZ"),
                    "end":         _local_to_utc(slot.end,   tz).strftime("%Y-%m-%dT%H:%M:%SZ"),
                    "date":        day.isoformat(),
                    "localStart":  slot.start.strftime("%H:%M"),
                    "localEnd":    slot.end.strftime("%H:%M"),
                    "ticketKey":   "",
                    "hours":       round(slot.duration_hours, 4),
                    "isExisting":  True,
                })
            else:
                plan.append({
                    "description": slot.description,
                    "projectId":   slot.project_id,
                    "start":       _local_to_utc(slot.start, tz).strftime("%Y-%m-%dT%H:%M:%SZ"),
                    "end":         _local_to_utc(slot.end,   tz).strftime("%Y-%m-%dT%H:%M:%SZ"),
                    "date":        day.isoformat(),
                    "localStart":  slot.start.strftime("%H:%M"),
                    "localEnd":    slot.end.strftime("%H:%M"),
                    "ticketKey":   slot.description.split()[0] if slot.description else "",
                    "hours":       round(slot.duration_hours, 4),
                    "isExisting":  False,
                })
    return plan


def _write_plan_json(path: str, plan: list[dict]) -> None:
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(plan, fh, indent=2)
    print(f"Plan written to: {path}")


def _write_plan_csv(path: str, plan: list[dict]) -> None:
    with open(path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(["date", "start", "end", "ticket", "description", "project_id", "hours"])
        for e in plan:
            desc_parts = e["description"].split(None, 1)
            ticket = desc_parts[0] if not e["isExisting"] and len(desc_parts) > 1 else ""
            title  = desc_parts[1] if not e["isExisting"] and len(desc_parts) > 1 else e["description"]
            writer.writerow([
                e["date"], e["localStart"], e["localEnd"],
                ticket, title, e["projectId"], e["hours"],
            ])
    print(f"Plan CSV written to: {path}")


def _push_from_json(plan_path: str, api_key: str, workspace_id: str, tz: ZoneInfo) -> None:
    """Read plan.json and submit only non-existing entries to Clockify."""
    with open(plan_path, encoding="utf-8") as fh:
        plan: list[dict] = json.load(fh)

    to_submit = [e for e in plan if not e.get("isExisting")]
    if not to_submit:
        print("No new entries found in plan.json.")
        return

    print(f"Submitting {len(to_submit)} entries from {plan_path} …\n")
    session = requests.Session()
    session.headers.update({"X-Api-Key": api_key, "Content-Type": "application/json"})

    created = failed = 0
    for e in to_submit:
        label = f"{e['description']} | {e['date']} {e['localStart']}–{e['localEnd']}"
        payload = {
            "description": e["description"],
            "projectId":   e["projectId"],
            "start":       e["start"],
            "end":         e["end"],
        }
        try:
            resp = session.post(
                f"{CLOCKIFY_API_BASE}/workspaces/{workspace_id}/time-entries",
                json=payload,
            )
            if resp.status_code in (200, 201):
                print(f"{GREEN}✓{RESET} {label}")
                created += 1
            else:
                try:
                    msg = resp.json().get("message", resp.text)
                except Exception:
                    msg = resp.text
                print(f"{RED}✗{RESET} {label} | HTTP {resp.status_code}: {msg}")
                failed += 1
        except requests.RequestException as exc:
            print(f"{RED}✗{RESET} {label} | Error: {exc}")
            failed += 1
        time.sleep(RATE_LIMIT_DELAY_S)

    print(f"\nCreated: {created} entries | Failed: {failed}")


# ─── CLI entry point ──────────────────────────────────────────────────────────

def _parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(
        description="Fill Clockify with ticket time entries for a given month.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=(
            "Examples:\n"
            "  python clockify_fill.py --dry-run\n"
            "  python clockify_fill.py --month 2026-04 --dry-run\n"
            "  python clockify_fill.py --from-json plan.json --config config.json\n"
        ),
    )
    p.add_argument("--tickets",      default="tickets.txt",        metavar="PATH",
                   help="Path to tickets file (default: tickets.txt)")
    p.add_argument("--report",       default="clockify_report.csv", metavar="PATH",
                   help="Path to Clockify CSV export (default: clockify_report.csv)")
    p.add_argument("--config",       default="config.json",        metavar="PATH",
                   help="Path to config file (default: config.json)")
    p.add_argument("--dry-run",      action="store_true",
                   help="Show plan only, skip API calls")
    p.add_argument("--month",        default=None, metavar="YYYY-MM",
                   help="Override 'month' from config")
    p.add_argument("--output-json",  default=None, metavar="PATH",
                   help="Write plan as JSON to this path (used with --dry-run)")
    p.add_argument("--output-csv",   default=None, metavar="PATH",
                   help="Write plan as CSV to this path (used with --dry-run)")
    p.add_argument("--from-json",    default=None, metavar="PATH",
                   help="Submit entries from a saved plan.json (skips generation)")
    p.add_argument("--invoice",      action="store_true",
                   help="Generate .xlsx invoice from plan.json (requires openpyxl)")
    p.add_argument("--plan",         default=None, metavar="PATH",
                   help="Path to plan.json (used with --invoice; default: same dir as config)")
    p.add_argument("--invoice-output", default=None, metavar="PATH",
                   help="Output .xlsx path (used with --invoice; default: same dir as config)")
    p.add_argument("--skip-days",    default=None, metavar="DATES",
                   help="Comma-separated YYYY-MM-DD dates to exclude from scheduling (e.g. vacation)")
    return p.parse_args()


def main():
    args = _parse_args()

    # ── --invoice: generate .xlsx invoice from plan.json ───────────────────────
    if args.invoice:
        try:
            from openpyxl import Workbook            # type: ignore
            from openpyxl.styles import Font, Alignment, Border, Side, PatternFill  # type: ignore
            from openpyxl.worksheet.page import PageMargins  # type: ignore
        except ImportError:
            print("Error: 'openpyxl' not found. Install with: pip install openpyxl")
            sys.exit(1)

        from pathlib import Path

        config_path = Path(args.config)
        try:
            config = load_config(str(config_path))
        except FileNotFoundError:
            print(f"Error: config file not found at '{config_path}'")
            sys.exit(1)

        if "monthly_total" not in config and "hourly_rate" not in config:
            print("Error: 'monthly_total' not set in config.json. Add e.g. \"monthly_total\": 3500")
            sys.exit(1)

        month_str = args.month or config.get("month", "")
        if not month_str:
            print("Error: month not specified — use --month YYYY-MM or set 'month' in config.json")
            sys.exit(1)
        config["month"] = month_str

        plan_path = Path(args.plan) if args.plan else config_path.parent / "plan.json"
        try:
            plan: list[dict] = json.loads(plan_path.read_text(encoding="utf-8"))
        except FileNotFoundError:
            print(f"Error: plan.json not found at '{plan_path}'")
            sys.exit(1)

        if not plan:
            print("Error: plan.json is empty — generate a plan first.")
            sys.exit(1)

        # Compute output path
        invoice_number = str(config.get("invoice_number", "espana"))[:10]
        if args.invoice_output:
            output_path = Path(args.invoice_output)
        else:
            now = datetime.now()
            base = f"invoice_{invoice_number}_{now.year}{now.month:02d}{now.day:02d}"
            output_path = config_path.parent / f"{base}.xlsx"
            # Avoid overwriting a file that may be open — append _N suffix
            counter = 1
            while output_path.exists():
                output_path = config_path.parent / f"{base}_{counter}.xlsx"
                counter += 1

        def calc_row_height(text: str, col_width: float, font_size: float = 10.0) -> float:
            col_px = col_width * 7
            avg_char_px = font_size * 0.55
            chars_per_line = max(1, int(col_px / avg_char_px))
            lines = 0
            for paragraph in str(text).split("\n"):
                if len(paragraph) == 0:
                    lines += 1
                else:
                    lines += max(1, -(-len(paragraph) // chars_per_line))
            return max(15.75, lines * font_size * 1.5)

        # ── Build workbook ────────────────────────────────────────────────────
        year_int, month_int = map(int, month_str.split("-"))
        if month_int == 12:
            inv_last = date(year_int + 1, 1, 1) - timedelta(days=1)
        else:
            inv_last = date(year_int, month_int + 1, 1) - timedelta(days=1)
        while inv_last.weekday() >= 5:
            inv_last -= timedelta(days=1)

        monthly_total = float(config.get("monthly_total", config.get("hourly_rate", 0)))

        entries_by_date: dict[str, list] = {}
        for e in plan:
            entries_by_date.setdefault(e["date"], []).append(e)
        sorted_entries = []
        for d in sorted(entries_by_date.keys(), reverse=True):
            sorted_entries.extend(entries_by_date[d])

        wb = Workbook()
        ws = wb.active

        ws.column_dimensions["A"].width = 16.71
        ws.column_dimensions["B"].width = 59.29
        ws.column_dimensions["C"].width = 14.86

        for r, h in [(1, 34.5), (2, 71.25), (5, 33.0), (8, 42.75), (9, 71.25),
                     (10, 15.0), (11, 15.0), (12, 18.0), (13, 21.0), (14, 20.25), (15, 21.75)]:
            ws.row_dimensions[r].height = h

        ws.merge_cells("A1:B1")
        ws.merge_cells("A12:B12")
        ws.merge_cells("A13:B13")
        ws.merge_cells("A14:B14")
        ws.merge_cells("A15:B15")

        thin        = Side(border_style="thin")
        all_borders = Border(top=thin, bottom=thin, left=thin, right=thin)
        lr_border   = Border(left=thin, right=thin)
        white_fill  = PatternFill(fill_type="solid", fgColor="FFFFFFFF")
        no_fill     = PatternFill(fill_type=None)
        cv          = Alignment(horizontal="center", vertical="center")

        arial11      = Font(name="Arial", size=11)
        arial11_bold = Font(name="Arial", size=11, bold=True)
        verdana      = Font(name="Verdana")

        # ── Header block (rows 1–15) ───────────────────────────────────────────
        ws["A1"].value     = "SENDER_NAME"
        ws["A1"].font      = Font(name="Arial", size=18, bold=True)

        ws["A2"].value     = "EBAN"
        ws["A2"].font      = arial11
        ws["A2"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        ws["B2"].value     = "SENDER_IBAN"
        ws["B2"].font      = arial11
        ws["B2"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        ws["A3"].value     = "BIC"
        ws["A3"].font      = arial11
        ws["A3"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        ws["B3"].value     = "SENDER_BIC"
        ws["B3"].font      = arial11
        ws["B3"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        ws["A4"].value     = "Bank name"
        ws["A4"].font      = arial11
        ws["A4"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        ws["B4"].value     = "SENDER_BANK_NAME"
        ws["B4"].font      = arial11
        ws["B4"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        ws["A5"].value     = "Address "
        ws["A5"].font      = arial11
        ws["A5"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        ws["B5"].value     = "SENDER_BANK_ADDRESS"
        ws["B5"].font      = arial11
        ws["B5"].alignment = Alignment(horizontal="left", vertical="center")

        ws["A7"].value     = "Address:"
        ws["A7"].font      = arial11_bold
        ws["A7"].alignment = Alignment(horizontal="left", wrap_text=True)

        ws["B7"].value     = "DATE:"
        ws["B7"].font      = arial11_bold
        ws["B7"].alignment = Alignment(horizontal="right")

        now = datetime.now()
        ws["C7"].value        = datetime(now.year, now.month, now.day)
        ws["C7"].number_format = "mm-dd-yy"

        ws["A8"].value     = "SENDER_ADDRESS1"
        ws["A8"].font      = arial11
        ws["A8"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        ws["B8"].value     = "INVOICE #"
        ws["B8"].font      = arial11_bold
        ws["B8"].alignment = Alignment(horizontal="right", vertical="center")

        ws["C8"].value     = invoice_number
        ws["C8"].font      = arial11
        ws["C8"].alignment = Alignment(horizontal="left", vertical="center")

        ws["A9"].value     = "SENDER_ADDRESS2"
        ws["A9"].font      = arial11
        ws["A9"].alignment = Alignment(vertical="center", wrap_text=True)

        ws["B9"].value     = "FOR:"
        ws["B9"].font      = arial11_bold
        ws["B9"].alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)

        ws["C9"].value     = "Software Development"
        ws["C9"].font      = arial11
        ws["C9"].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

        ws["A10"].value        = "SENDER_PHONE"
        ws["A10"].font         = arial11
        ws["A10"].alignment    = Alignment(vertical="center")
        ws["A10"].number_format = "@"

        ws["A11"].value     = "Bill to:"
        ws["A11"].font      = arial11_bold
        ws["A11"].alignment = Alignment(horizontal="left")

        ws["A12"].value     = "BILLTO_NAME"
        ws["A12"].font      = arial11
        ws["A12"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        ws["A13"].value     = "BILLTO_ADDRESS1"
        ws["A13"].font      = arial11
        ws["A13"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        ws["A14"].value     = "BILLTO_ADDRESS2"
        ws["A14"].font      = arial11
        ws["A14"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        ws["A15"].value        = "BILLTO_PHONE"
        ws["A15"].font         = arial11
        ws["A15"].alignment    = Alignment(horizontal="left", vertical="top")
        ws["A15"].number_format = "@"

        # ── Table header (row 16) ──────────────────────────────────────────────
        for col, label in [("A", "Date"), ("B", "DESCRIPTION"), ("C", "HOURS")]:
            cell = ws[f"{col}16"]
            cell.value     = label
            cell.alignment = cv
            cell.border    = all_borders
            cell.fill      = white_fill

        # ── Time entries (rows 17+) ────────────────────────────────────────────
        row = 17
        for i, entry in enumerate(sorted_entries):
            entry_date = datetime.strptime(entry["date"], "%Y-%m-%d").date()
            fill = white_fill if i % 2 == 1 else no_fill

            cell_a = ws.cell(row=row, column=1)
            cell_a.value         = entry_date
            cell_a.number_format = r"dd\/MM\/yyyy"
            cell_a.alignment     = cv
            cell_a.border        = lr_border
            cell_a.fill          = fill
            cell_a.font          = verdana

            desc = entry.get("description", "")
            key  = entry.get("ticketKey", "")
            if not entry.get("isExisting", False) and key and desc.startswith(key):
                desc = f"{key}   {desc[len(key):].lstrip()}"

            cell_b = ws.cell(row=row, column=2)
            cell_b.value     = desc
            cell_b.alignment = Alignment(vertical="top", wrap_text=True)
            cell_b.border    = lr_border
            cell_b.fill      = fill
            cell_b.font      = verdana

            cell_c = ws.cell(row=row, column=3)
            total_sec = int(entry.get("hours", 0) * 3600)
            hh = total_sec // 3600; mm_v = (total_sec % 3600) // 60; ss = total_sec % 60
            cell_c.value         = f"{hh:02d}:{mm_v:02d}:{ss:02d}"
            cell_c.number_format = "HH:mm:ss"
            cell_c.alignment     = cv
            cell_c.border        = lr_border
            cell_c.fill          = fill
            cell_c.font          = verdana

            ws.row_dimensions[row].height = calc_row_height(desc, 59.29)
            row += 1

        total_row = row

        # ── TOTAL row ──────────────────────────────────────────────────────────
        ws.cell(row=total_row, column=2).value        = "TOTAL"
        ws.cell(row=total_row, column=2).font         = arial11_bold
        ws.cell(row=total_row, column=2).alignment    = Alignment(horizontal="right")
        ws.cell(row=total_row, column=2).border       = all_borders
        ws.cell(row=total_row, column=2).fill         = white_fill

        ws.cell(row=total_row, column=3).value        = float(monthly_total)
        ws.cell(row=total_row, column=3).font         = arial11_bold
        ws.cell(row=total_row, column=3).alignment    = Alignment(horizontal="right")
        ws.cell(row=total_row, column=3).border       = all_borders
        ws.cell(row=total_row, column=3).fill         = white_fill
        ws.cell(row=total_row, column=3).number_format = '"$"#,##0_);[Red]("$"#,##0)'

        ws.row_dimensions[total_row].height = 15.75

        # ── Footer ─────────────────────────────────────────────────────────────
        footer_lines = [
            "Make all checks payable to SENDER_NAME",
            "If you have any questions concerning this invoice, use the following contact information:",
            "SENDER_NAME, SENDER_PHONE, SENDER_EMAIL",
            "THANK YOU FOR YOUR BUSINESS! ",
        ]
        footer_start = total_row + 1
        for i, line in enumerate(footer_lines):
            r = footer_start + i
            ws.merge_cells(f"A{r}:C{r}")
            ws[f"A{r}"].value     = line
            ws[f"A{r}"].alignment = Alignment(horizontal="center", wrap_text=True)
            ws[f"A{r}"].font      = arial11
            ws.row_dimensions[r].height = 15.75

        # Last footer line: bold + vertical bottom
        thank_you_row = footer_start + len(footer_lines) - 1
        ws[f"A{thank_you_row}"].font      = arial11_bold
        ws[f"A{thank_you_row}"].alignment = Alignment(horizontal="center", vertical="bottom")

        # ── Page setup ─────────────────────────────────────────────────────────
        last_row = thank_you_row
        ws.print_area = f"A1:C{last_row}"
        ws.page_setup.fitToPage   = True
        ws.page_setup.fitToWidth  = 1
        ws.page_setup.fitToHeight = 0
        ws.page_setup.orientation = "portrait"
        ws.page_setup.paperSize   = ws.PAPERSIZE_A4
        ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.75, bottom=0.75, header=0.3, footer=0.3)
        ws.sheet_properties.pageSetUpPr.fitToPage = True

        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(output_path))
        print(f"Invoice written to: {output_path}")
        sys.exit(0)

    # ── --from-json: submit a pre-generated plan without regenerating ──────────
    if args.from_json:
        try:
            cfg = load_config(args.config)
        except FileNotFoundError:
            print(f"Error: config file not found at '{args.config}'")
            sys.exit(1)
        api_key      = cfg.get("clockify_api_key", "")
        workspace_id = cfg.get("workspace_id",     "")
        tz_name      = cfg.get("timezone", "UTC")
        if not api_key or api_key.startswith("YOUR_"):
            print("Error: 'clockify_api_key' is not configured.")
            sys.exit(1)
        if not workspace_id or workspace_id.startswith("YOUR_"):
            print("Error: 'workspace_id' is not configured.")
            sys.exit(1)
        try:
            tz = ZoneInfo(tz_name)
        except Exception:
            print(f"Error: unknown timezone '{tz_name}'")
            sys.exit(1)
        _push_from_json(args.from_json, api_key, workspace_id, tz)
        return

    # ── Config ────────────────────────────────────────────────────────────────
    try:
        cfg = load_config(args.config)
    except FileNotFoundError:
        print(f"Error: config file not found at '{args.config}'")
        sys.exit(1)
    except json.JSONDecodeError as exc:
        print(f"Error: invalid JSON in config: {exc}")
        sys.exit(1)

    target_month = args.month or cfg.get("month", "")
    if not target_month:
        print("Error: 'month' not set in config and --month not provided.")
        sys.exit(1)
    try:
        datetime.strptime(target_month, "%Y-%m")
    except ValueError:
        print(f"Error: month must be in YYYY-MM format, got '{target_month}'")
        sys.exit(1)

    work_start   = cfg.get("work_start",    "12:00")
    work_end     = cfg.get("work_end",      "20:00")
    skip_weekends = cfg.get("skip_weekends", True)
    tz_name      = cfg.get("timezone",      "UTC")
    api_key      = cfg.get("clockify_api_key", "")
    workspace_id = cfg.get("workspace_id",  "")

    try:
        tz = ZoneInfo(tz_name)
    except Exception:
        print(f"Error: unknown timezone '{tz_name}'")
        sys.exit(1)

    # ── Tickets ───────────────────────────────────────────────────────────────
    try:
        tickets = load_tickets(args.tickets)
    except FileNotFoundError:
        print(f"Error: tickets file not found at '{args.tickets}'")
        sys.exit(1)

    if not tickets:
        print("Error: no tickets found in tickets file.")
        sys.exit(1)
    print(f"Loaded {len(tickets)} ticket(s) from '{args.tickets}'.")

    # ── CSV report ────────────────────────────────────────────────────────────
    existing_by_day = parse_clockify_csv(args.report, target_month, work_start, work_end)
    print(f"Existing entries found on {len(existing_by_day)} day(s).")

    # ── Working days ──────────────────────────────────────────────────────────
    working_days = get_working_days(target_month, skip_weekends)
    if args.skip_days:
        skip_set = {s.strip() for s in args.skip_days.split(",") if s.strip()}
        working_days = [d for d in working_days if d.isoformat() not in skip_set]
    print(f"Working days in {target_month}: {len(working_days)}")

    # ── Distribute ────────────────────────────────────────────────────────────
    new_entries = distribute_tickets(working_days, existing_by_day, tickets, work_start, work_end)

    if not new_entries:
        print("No free slots available to fill — all days may already be at 8h.")
        if args.output_json or args.output_csv:
            plan_data = _build_plan_json(working_days, existing_by_day, [], tz)
            if args.output_json:
                _write_plan_json(args.output_json, plan_data)
            if args.output_csv:
                _write_plan_csv(args.output_csv, plan_data)
        sys.exit(0)

    # ── Display plan ──────────────────────────────────────────────────────────
    print_plan(working_days, existing_by_day, new_entries)

    # ── Write output files if requested ───────────────────────────────────────
    if args.output_json or args.output_csv:
        plan_data = _build_plan_json(working_days, existing_by_day, new_entries, tz)
        if args.output_json:
            _write_plan_json(args.output_json, plan_data)
        if args.output_csv:
            _write_plan_csv(args.output_csv, plan_data)

    if args.dry_run:
        print("Dry-run mode — no entries will be created.")
        sys.exit(0)

    # ── Validate API credentials before asking ─────────────────────────────────
    if not api_key or api_key.startswith("YOUR_"):
        print("Error: 'clockify_api_key' is not configured.\n"
              "       Edit config.json or run with --dry-run to preview only.")
        sys.exit(1)
    if not workspace_id or workspace_id.startswith("YOUR_"):
        print("Error: 'workspace_id' is not configured.\n"
              "       Edit config.json or run with --dry-run to preview only.")
        sys.exit(1)

    # ── Approval ──────────────────────────────────────────────────────────────
    try:
        answer = input("Proceed with creating these entries in Clockify? (yes/no): ").strip().lower()
    except (EOFError, KeyboardInterrupt):
        print("\nAborted. No entries created.")
        sys.exit(0)

    if answer not in ("yes", "y"):
        print("Aborted. No entries created.")
        sys.exit(0)

    # ── Push to API ───────────────────────────────────────────────────────────
    push_entries(new_entries, api_key, workspace_id, tz)


if __name__ == "__main__":
    main()
