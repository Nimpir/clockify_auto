#!/usr/bin/env python3
"""
generate_invoice.py — Generate an .xlsx invoice from plan.json

Usage:
    python generate_invoice.py --month 2026-03
    python generate_invoice.py --month 2026-03 --config path/to/config.json
"""

import argparse
import json
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.worksheet.page import PageMargins
except ImportError:
    print("Error: 'openpyxl' not found. Install with: pip install openpyxl")
    sys.exit(1)

# Ensure Unicode output on Windows
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")


# ─── Helpers ──────────────────────────────────────────────────────────────────

def _strip_json_comments(text: str) -> str:
    """Remove // line-comments that are not inside JSON string literals."""
    result, in_string, i = [], False, 0
    while i < len(text):
        ch = text[i]
        if ch == '\\' and in_string:
            result.append(ch); i += 1
            if i < len(text):
                result.append(text[i]); i += 1
            continue
        if ch == '"':
            in_string = not in_string
            result.append(ch); i += 1
            continue
        if not in_string and text[i:i+2] == '//':
            while i < len(text) and text[i] != '\n':
                i += 1
            continue
        result.append(ch); i += 1
    return ''.join(result)


def last_working_day(year: int, month: int) -> date:
    """Return the last Monday–Friday day of the given month."""
    if month == 12:
        last = date(year + 1, 1, 1) - timedelta(days=1)
    else:
        last = date(year, month + 1, 1) - timedelta(days=1)
    while last.weekday() >= 5:
        last -= timedelta(days=1)
    return last


def format_description(entry: dict) -> str:
    """Return invoice description string for a plan entry."""
    desc = entry.get("description", "")
    if entry.get("isExisting", False):
        return desc
    key = entry.get("ticketKey", "")
    if key and desc.startswith(key):
        rest = desc[len(key):].lstrip()
        return f"{key}   {rest}"
    return desc


def calc_row_height(text: str, col_width: float, font_size: float = 10.0) -> float:
    """Estimate required row height in points to fit wrapped text in a cell."""
    col_px = col_width * 7          # Excel char unit ≈ 7px
    avg_char_px = font_size * 0.55  # Verdana is slightly wider
    chars_per_line = max(1, int(col_px / avg_char_px))
    lines = 0
    for paragraph in str(text).split("\n"):
        if len(paragraph) == 0:
            lines += 1
        else:
            lines += max(1, -(-len(paragraph) // chars_per_line))  # ceiling div
    return max(15.75, lines * font_size * 1.5)


# ─── Invoice builder ──────────────────────────────────────────────────────────

def generate_invoice(plan: list[dict], config: dict, output_path: Path, table_only: bool = False) -> None:
    year_int, month_int = map(int, config["month"].split("-"))
    inv_date      = last_working_day(year_int, month_int)
    monthly_total = float(config.get("monthly_total", config.get("hourly_rate", 0)))

    # Sort: newest date first, preserve original order within each day
    entries_by_date: dict[str, list[dict]] = {}
    for e in plan:
        entries_by_date.setdefault(e["date"], []).append(e)
    sorted_entries: list[dict] = []
    for d in sorted(entries_by_date.keys(), reverse=True):
        sorted_entries.extend(entries_by_date[d])

    wb = Workbook()
    ws = wb.active

    # ── Column widths ──────────────────────────────────────────────────────────
    ws.column_dimensions["A"].width = 16.71
    ws.column_dimensions["B"].width = 59.29
    ws.column_dimensions["C"].width = 14.86

    # ── Shared style objects ───────────────────────────────────────────────────
    thin        = Side(border_style="thin")
    all_borders = Border(top=thin, bottom=thin, left=thin, right=thin)
    lr_border   = Border(left=thin, right=thin)
    white_fill  = PatternFill(fill_type="solid", fgColor="FFFFFFFF")
    no_fill     = PatternFill(fill_type=None)
    cv          = Alignment(horizontal="center", vertical="center")

    arial11      = Font(name="Arial", size=11)
    arial11_bold = Font(name="Arial", size=11, bold=True)
    verdana      = Font(name="Verdana")

    if not table_only:
        # ── Row heights (header block) ─────────────────────────────────────────
        for r, h in [(1, 34.5), (2, 71.25), (5, 33.0), (8, 42.75), (9, 71.25),
                     (10, 15.0), (11, 15.0), (12, 18.0), (13, 21.0), (14, 20.25), (15, 21.75)]:
            ws.row_dimensions[r].height = h

        # ── Merged cells (header) ─────────────────────────────────────────────
        ws.merge_cells("A1:B1")
        ws.merge_cells("A12:B12")
        ws.merge_cells("A13:B13")
        ws.merge_cells("A14:B14")
        ws.merge_cells("A15:B15")

        # ── Header block (rows 1–15) ───────────────────────────────────────────
        sender   = config.get("sender", {})
        bill_to  = config.get("bill_to", {})

        ws["A1"].value     = sender.get("name", "")
        ws["A1"].font      = Font(name="Arial", size=18, bold=True)

        ws["A2"].value     = "IBAN"
        ws["A2"].font      = arial11
        ws["A2"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        ws["B2"].value     = sender.get("iban", "")
        ws["B2"].font      = arial11
        ws["B2"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        ws["A3"].value     = "BIC"
        ws["A3"].font      = arial11
        ws["A3"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        ws["B3"].value     = sender.get("bic", "")
        ws["B3"].font      = arial11
        ws["B3"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        ws["A4"].value     = "Bank name"
        ws["A4"].font      = arial11
        ws["A4"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        ws["B4"].value     = sender.get("bank_name", "")
        ws["B4"].font      = arial11
        ws["B4"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        ws["A5"].value     = "Address"
        ws["A5"].font      = arial11
        ws["A5"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        ws["B5"].value     = sender.get("bank_address", "")
        ws["B5"].font      = arial11
        ws["B5"].alignment = Alignment(horizontal="left", vertical="center")

        ws["A7"].value     = "Address:"
        ws["A7"].font      = arial11_bold
        ws["A7"].alignment = Alignment(horizontal="left", wrap_text=True)

        ws["B7"].value     = "DATE:"
        ws["B7"].font      = arial11_bold
        ws["B7"].alignment = Alignment(horizontal="right")

        today = datetime.now()
        ws["C7"].value        = datetime(today.year, today.month, today.day)
        ws["C7"].number_format = "mm-dd-yy"

        ws["A8"].value     = sender.get("address1", "")
        ws["A8"].font      = arial11
        ws["A8"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        ws["B8"].value     = "INVOICE #"
        ws["B8"].font      = arial11_bold
        ws["B8"].alignment = Alignment(horizontal="right", vertical="center")

        invoice_number = str(config.get("invoice_number", ""))[:10]
        ws["C8"].value     = invoice_number
        ws["C8"].font      = arial11
        ws["C8"].alignment = Alignment(horizontal="left", vertical="center")

        ws["A9"].value     = sender.get("address2", "")
        ws["A9"].font      = arial11
        ws["A9"].alignment = Alignment(vertical="center", wrap_text=True)

        ws["B9"].value     = "FOR:"
        ws["B9"].font      = arial11_bold
        ws["B9"].alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)

        ws["C9"].value     = config.get("service", "Software Development")
        ws["C9"].font      = arial11
        ws["C9"].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

        ws["A10"].value        = sender.get("phone", "")
        ws["A10"].font         = arial11
        ws["A10"].alignment    = Alignment(vertical="center")
        ws["A10"].number_format = "@"

        ws["A11"].value     = "Bill to:"
        ws["A11"].font      = arial11_bold
        ws["A11"].alignment = Alignment(horizontal="left")

        ws["A12"].value     = bill_to.get("name", "")
        ws["A12"].font      = arial11
        ws["A12"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        ws["A13"].value     = bill_to.get("address1", "")
        ws["A13"].font      = arial11
        ws["A13"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        ws["A14"].value     = bill_to.get("address2", "")
        ws["A14"].font      = arial11
        ws["A14"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        ws["A15"].value        = bill_to.get("phone", "")
        ws["A15"].font         = arial11
        ws["A15"].alignment    = Alignment(horizontal="left", vertical="top")
        ws["A15"].number_format = "@"

    # ── Table header ───────────────────────────────────────────────────────────
    table_header_row = 1 if table_only else 16
    for col, label in [("A", "Date"), ("B", "DESCRIPTION"), ("C", "HOURS")]:
        cell = ws[f"{col}{table_header_row}"]
        cell.value     = label
        cell.alignment = cv
        cell.border    = all_borders
        cell.fill      = white_fill

    # ── Time entries ───────────────────────────────────────────────────────────
    row = table_header_row + 1
    for i, entry in enumerate(sorted_entries):
        entry_date = datetime.strptime(entry["date"], "%Y-%m-%d").date()
        fill = white_fill if i % 2 == 1 else no_fill

        cell_a = ws.cell(row=row, column=1)
        cell_a.value         = entry_date
        cell_a.number_format = r"dd\/MM\/yyyy"
        cell_a.alignment     = cv
        cell_a.border        = lr_border
        cell_a.fill          = fill
        cell_a.font          = verdana

        desc = format_description(entry)
        cell_b = ws.cell(row=row, column=2)
        cell_b.value     = desc
        cell_b.alignment = Alignment(vertical="top", wrap_text=True)
        cell_b.border    = lr_border
        cell_b.fill      = fill
        cell_b.font      = verdana

        cell_c = ws.cell(row=row, column=3)
        total_sec = int(entry.get("hours", 0) * 3600)
        hh = total_sec // 3600
        mm = (total_sec % 3600) // 60
        ss = total_sec % 60
        cell_c.value         = f"{hh:02d}:{mm:02d}:{ss:02d}"
        cell_c.number_format = "HH:mm:ss"
        cell_c.alignment     = cv
        cell_c.border        = lr_border
        cell_c.fill          = fill
        cell_c.font          = verdana

        ws.row_dimensions[row].height = calc_row_height(desc, 59.29)
        row += 1

    last_data_row = row - 1
    total_row = row

    # ── TOTAL row ──────────────────────────────────────────────────────────────
    ws.cell(row=total_row, column=2).value        = "TOTAL"
    ws.cell(row=total_row, column=2).font         = arial11_bold
    ws.cell(row=total_row, column=2).alignment    = Alignment(horizontal="right")
    ws.cell(row=total_row, column=2).border       = all_borders
    ws.cell(row=total_row, column=2).fill         = white_fill

    ws.cell(row=total_row, column=3).value        = float(monthly_total)
    ws.cell(row=total_row, column=3).font         = arial11_bold
    ws.cell(row=total_row, column=3).alignment    = Alignment(horizontal="right")
    ws.cell(row=total_row, column=3).border       = all_borders
    ws.cell(row=total_row, column=3).fill         = white_fill
    ws.cell(row=total_row, column=3).number_format = '"$"#,##0_);[Red]("$"#,##0)'

    ws.row_dimensions[total_row].height = 15.75

    # ── Footer ─────────────────────────────────────────────────────────────────
    last_row = total_row
    if not table_only:
        sender = config.get("sender", {})
        sender_name  = sender.get("name", "")
        sender_phone = sender.get("phone", "")
        sender_email = sender.get("email", "")
        contact_line = ", ".join(x for x in [sender_name, sender_phone, sender_email] if x)
        footer_lines = [
            f"Make all checks payable to {sender_name}".strip(),
            "If you have any questions concerning this invoice, use the following contact information:",
            contact_line,
            "THANK YOU FOR YOUR BUSINESS! ",
        ]
        footer_start = total_row + 1
        for i, line in enumerate(footer_lines):
            r = footer_start + i
            ws.merge_cells(f"A{r}:C{r}")
            ws[f"A{r}"].value     = line
            ws[f"A{r}"].alignment = Alignment(horizontal="center", wrap_text=True)
            ws[f"A{r}"].font      = arial11
            ws.row_dimensions[r].height = 15.75

        # Last footer line bold
        thank_you_row = footer_start + len(footer_lines) - 1
        ws[f"A{thank_you_row}"].font      = arial11_bold
        ws[f"A{thank_you_row}"].alignment = Alignment(horizontal="center", vertical="bottom")
        last_row = thank_you_row

    # ── Page setup ─────────────────────────────────────────────────────────────
    ws.print_area = f"A1:C{last_row}"
    ws.page_setup.fitToPage   = True
    ws.page_setup.fitToWidth  = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize   = ws.PAPERSIZE_A4
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.75, bottom=0.75, header=0.3, footer=0.3)
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"Invoice written to: {output_path}")


# ─── CLI ──────────────────────────────────────────────────────────────────────

def _parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(
        description="Generate .xlsx invoice from plan.json",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=(
            "Examples:\n"
            "  python generate_invoice.py --month 2026-03\n"
            "  python generate_invoice.py --month 2026-03 --config docs/2026-03/config.json\n"
        ),
    )
    p.add_argument("--month",  required=True, metavar="YYYY-MM", help="Target month")
    p.add_argument("--config", default="config.json", metavar="PATH",
                   help="Path to config.json (default: config.json)")
    p.add_argument("--plan",   default=None, metavar="PATH",
                   help="Path to plan.json (default: same directory as config)")
    p.add_argument("--output", default=None, metavar="PATH",
                   help="Output .xlsx path (default: same directory as config)")
    p.add_argument("--table-only", action="store_true",
                   help="Output only the time table — no header or footer (removes sensitive info)")
    return p.parse_args()


def main() -> None:
    args = _parse_args()

    # ── Config ────────────────────────────────────────────────────────────────
    config_path = Path(args.config)
    try:
        raw = config_path.read_text(encoding="utf-8")
        config = json.loads(_strip_json_comments(raw))
    except FileNotFoundError:
        print(f"Error: config file not found at '{config_path}'")
        sys.exit(1)
    except json.JSONDecodeError as exc:
        print(f"Error: invalid JSON in config: {exc}")
        sys.exit(1)

    config["month"] = args.month

    if "monthly_total" not in config and "hourly_rate" not in config:
        print("Error: 'monthly_total' not set in config.json. Add e.g. \"monthly_total\": 3500")
        sys.exit(1)

    # ── Plan ──────────────────────────────────────────────────────────────────
    plan_path = Path(args.plan) if args.plan else config_path.parent / "plan.json"
    try:
        plan: list[dict] = json.loads(plan_path.read_text(encoding="utf-8"))
    except FileNotFoundError:
        print(f"Error: plan.json not found at '{plan_path}'")
        sys.exit(1)

    if not plan:
        print("Error: plan.json is empty — generate a plan first.")
        sys.exit(1)

    # ── Output path ───────────────────────────────────────────────────────────
    if args.output:
        output_path = Path(args.output)
    else:
        inv_num = str(config.get("invoice_number", "espana"))[:10]
        now = datetime.now()
        base = f"invoice_{inv_num}_{now.year}{now.month:02d}{now.day:02d}"
        output_path = config_path.parent / f"{base}.xlsx"
        counter = 1
        while output_path.exists():
            output_path = config_path.parent / f"{base}_{counter}.xlsx"
            counter += 1

    generate_invoice(plan, config, output_path, table_only=args.table_only)


if __name__ == "__main__":
    main()
